from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import inspect
from sqlalchemy.orm import Session

from app.database import Base, get_db
from app.models.employee import Employee
from app.models.user import User
from app.routers.auth import require_roles

router = APIRouter(
    prefix="/api/reports",
    tags=["Reports and Export"],
)

HR_ROLES = ("ADMIN", "HR", "MANAGER")
PAYROLL_ROLES = ("ADMIN", "HR")

REPORT_MODELS = {
    "employees": ("employees", "Employee Directory"),
    "attendance": ("attendance", "Attendance Report"),
    "leaves": ("leave_requests", "Leave Report"),
    "payroll": ("payroll_records", "Payroll Report"),
    "attrition": ("attrition_predictions", "Attrition Prediction Report"),
}

SENSITIVE_FIELDS = {
    "hashed_password",
    "password",
    "password_hash",
    "secret",
    "token",
}


def get_model(table_name: str):
    """Find the mapped model for an existing database table."""
    for mapper in Base.registry.mappers:
        if mapper.local_table.name == table_name:
            return mapper.class_

    raise HTTPException(
        status_code=501,
        detail=(
            f"The model for table '{table_name}' is not registered. "
            "Check the model imports in app/main.py."
        ),
    )


def serialize_value(value):
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, (dict, list)):
        return str(value)

    return value


def get_columns(model):
    return [
        column.key
        for column in inspect(model).columns
        if column.key.lower() not in SENSITIVE_FIELDS
    ]


def get_report_data(db: Session, report_type: str, limit: int):
    if report_type not in REPORT_MODELS:
        raise HTTPException(
            status_code=404,
            detail=(
                "Unknown report type. Available reports: "
                + ", ".join(REPORT_MODELS.keys())
            ),
        )

    table_name, title = REPORT_MODELS[report_type]
    model = get_model(table_name)

    records = (
        db.query(model)
        .order_by(model.id.desc())
        .limit(limit)
        .all()
    )

    columns = get_columns(model)

    rows = []
    for record in records:
        rows.append([
            serialize_value(getattr(record, column))
            for column in columns
        ])

    return title, columns, rows


def create_excel(title, columns, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    sheet.append([title])
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(1, len(columns)),
    )

    sheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )
    sheet["A1"].fill = PatternFill(
        "solid",
        fgColor="17365D",
    )
    sheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    sheet.append(columns)

    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = sheet.dimensions

    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        maximum = max(
            (
                len(str(cell.value))
                for cell in column_cells
                if cell.value is not None
            ),
            default=10,
        )
        sheet.column_dimensions[column_letter].width = min(
            max(maximum + 2, 12),
            35,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_pdf(title, columns, rows):
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#17365D"),
    )

    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["BodyText"],
        fontSize=6,
        leading=7,
        wordWrap="CJK",
    )

    story = [
        Paragraph(title, title_style),
        Spacer(1, 5 * mm),
        Paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles["Normal"],
        ),
        Paragraph(f"Records included: {len(rows)}", styles["Normal"]),
        Spacer(1, 4 * mm),
    ]

    table_data = [
        [
            Paragraph(str(column).replace("_", " ").title(), cell_style)
            for column in columns
        ]
    ]

    for row in rows:
        table_data.append([
            Paragraph(
                str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
                cell_style,
            )
            for value in row
        ])

    if not columns:
        table_data = [["No columns available"]]

    available_width = landscape(A4)[0] - 20 * mm
    col_width = available_width / max(1, len(columns))

    table = Table(
        table_data,
        colWidths=[col_width] * max(1, len(columns)),
        repeatRows=1,
    )

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    story.append(table)
    document.build(story)
    output.seek(0)
    return output


@router.get("/")
def available_reports(
    current_user: User = Depends(require_roles(*HR_ROLES)),
):
    return {
        "reports": [
            {
                "type": report_type,
                "title": title,
                "formats": ["pdf", "xlsx"],
            }
            for report_type, (_, title) in REPORT_MODELS.items()
        ]
    }


@router.get("/{report_type}/excel")
def export_excel(
    report_type: str,
    limit: int = Query(default=5000, ge=1, le=10000),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(*HR_ROLES)),
):
    if report_type == "payroll" and current_user.role.upper() not in PAYROLL_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Only Admin and HR can export payroll reports.",
        )

    title, columns, rows = get_report_data(
        db,
        report_type,
        limit,
    )

    output = create_excel(title, columns, rows)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{report_type}_report.xlsx"'
        },
    )


@router.get("/{report_type}/pdf")
def export_pdf(
    report_type: str,
    limit: int = Query(default=1000, ge=1, le=3000),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(*HR_ROLES)),
):
    if report_type == "payroll" and current_user.role.upper() not in PAYROLL_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Only Admin and HR can export payroll reports.",
        )

    title, columns, rows = get_report_data(
        db,
        report_type,
        limit,
    )

    output = create_pdf(title, columns, rows)

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{report_type}_report.pdf"'
        },
    )
